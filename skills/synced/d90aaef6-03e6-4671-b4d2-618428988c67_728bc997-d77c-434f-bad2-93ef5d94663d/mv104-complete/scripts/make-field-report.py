#!/usr/bin/env python3
"""Build the two Word reports that accompany every MV-104 produced by this skill.

Usage:
  python3 make-field-report.py --worksheet worksheet.json --fill-report MV104_<id>.fill-report.json \
      --incident 2026001358 --out-dir ./out

worksheet.json layout (see SKILL.md "Worksheet"):
{
  "incident_id": "2026001358",
  "verbatim": {"<exact field>": {"value": "...", "source": "Origami p.1 'Incident Date'"}, ...},
  "inferred": {"<exact field>": {"value": "...", "why": "..."}, ...}
}

Outputs:
  MV104_<id>_Verbatim-Fields.docx  — Section 1: every value written to the PDF, in form
                                      order, labelled by the form's own field name and source.
                                      Section 2: REQUIRED fields still missing (bold + yellow).
  MV104_<id>_Inferred-Fields.docx  — Section 1: fields the skill believes it knows but did NOT
                                      write (form order, value only). Appendix A: the WHY for
                                      each, keyed by the same field. Section 2 repeats the
                                      missing-required list so either document stands alone.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

SKILL_ROOT: Path = Path(__file__).resolve().parent.parent
FIELD_MAP_PATH: Path = SKILL_ROOT / "references" / "mv104-field-map.json"
BODY_FONT: str = "Aptos"
BODY_SIZE_PT: int = 10
TABLE_STYLE: str = "Table Grid"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def form_order() -> dict[str, int]:
    return {d["field"]: d["order"] for d in load_json(FIELD_MAP_PATH)}


def new_document(title: str, subtitle: str) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(BODY_SIZE_PT)
    doc.add_heading(title, level=0)
    doc.add_paragraph(subtitle)
    return doc


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = TABLE_STYLE
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text


def add_missing_section(doc: Document, missing: list[dict]) -> None:
    doc.add_heading("Section 2 — REQUIRED fields MISSING (cannot be completed from the sources provided)", level=1)
    if not missing:
        doc.add_paragraph("None. Every required field was completed verbatim.")
        return
    for gap in missing:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(f"MISSING REQUIRED — {gap['label']}  [{gap['section']}]")
        run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def ordered_items(section: dict[str, dict], order: dict[str, int]) -> list[tuple[int, str, dict]]:
    items = [(order.get(field, 9999), field, data) for field, data in section.items()]
    return sorted(items, key=lambda item: item[0])


def build_verbatim_doc(ws: dict, missing: list[dict], out: Path) -> None:
    doc = new_document(f"MV-104 Field Report — Incident {ws['incident_id']}",
                       "Section 1 lists ONLY values transcribed verbatim from a source document; nothing here is inferred. "
                       "Fields are listed in the order they appear on form MV-104 (5/25), by the form's own field label.")
    doc.add_heading("Section 1 — Fields completed verbatim (form order)", level=1)
    rows = [[str(n), field, d["value"], d.get("source", "")] for n, field, d in ordered_items(ws.get("verbatim", {}), form_order())]
    add_table(doc, ["#", "MV-104 field (legal label)", "Value entered", "Source"], rows)
    add_missing_section(doc, missing)
    doc.save(out)


def build_inferred_doc(ws: dict, missing: list[dict], out: Path) -> None:
    doc = new_document(f"MV-104 Inferred Fields — Incident {ws['incident_id']}",
                       "These fields were NOT written to the PDF. They are values the skill believes are correct but that no "
                       "source states verbatim. Review each against Appendix A, then tell the skill which to accept.")
    items = ordered_items(ws.get("inferred", {}), form_order())
    doc.add_heading("Section 1 — Inferred fields and proposed values (form order)", level=1)
    add_table(doc, ["#", "MV-104 field (legal label)", "Proposed value"], [[str(n), f, d["value"]] for n, f, d in items])
    add_missing_section(doc, missing)
    doc.add_page_break()
    doc.add_heading("Appendix A — WHY each inferred value is proposed", level=1)
    add_table(doc, ["#", "MV-104 field (legal label)", "Why"], [[str(n), f, d["why"]] for n, f, d in items])
    doc.save(out)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worksheet", required=True, type=Path)
    parser.add_argument("--fill-report", required=True, type=Path)
    parser.add_argument("--incident", required=True)
    parser.add_argument("--out-dir", required=True, type=Path)
    args = parser.parse_args()
    ws = load_json(args.worksheet)
    ws["incident_id"] = args.incident
    missing: list[dict] = load_json(args.fill_report).get("missing_required", [])
    args.out_dir.mkdir(parents=True, exist_ok=True)
    verbatim_path = args.out_dir / f"MV104_{args.incident}_Verbatim-Fields.docx"
    inferred_path = args.out_dir / f"MV104_{args.incident}_Inferred-Fields.docx"
    build_verbatim_doc(ws, missing, verbatim_path)
    build_inferred_doc(ws, missing, inferred_path)
    print(f"wrote {verbatim_path}\nwrote {inferred_path}\nmissing_required={len(missing)}")


if __name__ == "__main__":
    main()
