#!/usr/bin/env python3
"""Preflight for the quarterly safety-bonus run. Run this FIRST and show the
output to Jeremy before any other script. It exits non-zero on any BLOCKER so
the pipeline cannot proceed on bad inputs by accident.

Checks
  Branch workbooks : one per expected branch, tabs present, header row sane,
                     DRIVER RECORD span found, ID column mostly numeric.
  HR term report   : sheet 'report', expected header, rows, term dates inside quarter.
  Origami export   : header found, required columns present (Preventable?,
                     Disciplinary action?, Employee Number, Loss Date, Location),
                     loss dates inside the quarter, every branch workbook has a
                     matching location in the export, blank employee numbers.

Usage:
  python3 validate_inputs.py --input-dir <branch files> --term-report <HR xlsx> \
      [--incidents <Origami xlsx>] --start 2026-07-01 --end 2026-09-30
"""
from __future__ import annotations

import argparse
import sys
import warnings
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from reconcile_safety_bonus import ACTIVE_TABS, TERMINATED_TAB, TERM_REPORT_HEADER_ROW, branch_from_filename, find_roster_columns, is_excluded_title  # noqa: E402

warnings.filterwarnings("ignore", category=UserWarning)

EXPECTED_BRANCHES: tuple[str, ...] = ("BB", "BNY", "BPA", "BDC", "FISH")
BRANCH_LOCATION_CODE: dict[str, str] = {"BB": "BMABXT", "BNY": "BNYBXT", "BPA": "BPABXT", "BDC": "BDCBXT", "FISH": "BFSBXT"}
REQUIRED_SHEET_HEADERS: tuple[str, ...] = ("NAME", "ID", "ELIG", "REASON", "ELIGIBLE FOR BONUS", "DRIVER RECORD")
REQUIRED_INCIDENT_COLUMNS: tuple[str, ...] = ("Occurrence Number", "Employee", "Employee Number", "Loss Date", "Location", "Preventable?", "Disciplinary action?", "Tier?")
TERM_REPORT_HEADERS: tuple[str, ...] = ("Employee Id", "First Name", "Last Name", "Sub Department", "Date Terminated")
MIN_NUMERIC_ID_SHARE: float = 0.9
BLANK_EMPLOYEE_NUMBER_WARN_SHARE: float = 0.05


class Findings:
    def __init__(self) -> None:
        self.blockers: list[str] = []
        self.warnings: list[str] = []
        self.info: list[str] = []

    def block(self, msg: str) -> None:
        self.blockers.append(msg)

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)

    def note(self, msg: str) -> None:
        self.info.append(msg)


def check_branch_workbooks(input_dir: Path, f: Findings) -> set[str]:
    found: dict[str, Path] = {}
    for path in sorted(input_dir.glob("*.xlsx")):
        branch = branch_from_filename(path)
        if branch:
            found[branch] = path
    for b in EXPECTED_BRANCHES:
        if b not in found:
            f.warn(f"No workbook for branch {b} (expected a file like {b}_Q#_YYYY_Safety_Bonus.xlsx). Branch will be skipped.")
    for branch, path in found.items():
        wb = openpyxl.load_workbook(path, read_only=True)
        for tab in (*ACTIVE_TABS, TERMINATED_TAB):
            if tab not in wb.sheetnames:
                f.block(f"{path.name}: tab '{tab}' is missing (has {wb.sheetnames}).")
                continue
            ws = wb[tab]
            rows = ws.iter_rows(values_only=True)
            header = [str(h).strip().upper() if h else "" for h in next(rows, [])]
            missing = [h for h in REQUIRED_SHEET_HEADERS if h not in header]
            if missing:
                f.block(f"{path.name} / {tab}: header row 1 is missing {missing}.")
                continue
            ids = [r[header.index("ID")] for r in rows if any(v not in (None, "") for v in r)]
            if ids and tab != TERMINATED_TAB:
                numeric = sum(1 for i in ids if isinstance(i, int)) / len(ids)
                if numeric < MIN_NUMERIC_ID_SHARE:
                    f.warn(f"{path.name} / {tab}: only {numeric:.0%} of ID cells are numbers - rows with text IDs (e.g. 'Vertex') can't be matched by ID.")
            f.note(f"{path.name} / {tab}: {len(ids)} driver rows")
    return set(found)


def check_term_report(path: Path, start: datetime, end: datetime, f: Findings) -> None:
    wb = openpyxl.load_workbook(path, read_only=True)
    if "report" not in wb.sheetnames:
        f.block(f"{path.name}: expected a sheet named 'report' (has {wb.sheetnames}).")
        return
    rows = list(wb["report"].iter_rows(values_only=True))
    header_idx = next((i for i, r in enumerate(rows) if r and "Employee Id" in r), None)
    if header_idx is None:
        f.block(f"{path.name}: could not find the 'Employee Id' header row (expected around row {TERM_REPORT_HEADER_ROW}).")
        return
    header = [str(h) for h in rows[header_idx]]
    missing = [h for h in TERM_REPORT_HEADERS if h not in header]
    if missing:
        f.block(f"{path.name}: missing columns {missing}.")
        return
    date_col = header.index("Date Terminated")
    terms = [r for r in rows[header_idx + 1:] if r and isinstance(r[0], int)]
    if not terms:
        f.block(f"{path.name}: no termination rows found.")
        return
    outside = [r for r in terms if isinstance(r[date_col], datetime) and not (start <= r[date_col] <= end)]
    f.note(f"{path.name}: {len(terms)} termed drivers, dates {min(r[date_col] for r in terms).date()} to {max(r[date_col] for r in terms).date()}")
    if outside:
        f.warn(f"{path.name}: {len(outside)} termination dates fall outside {start.date()}-{end.date()} - confirm the HR filter.")


def check_incidents(path: Path, start: datetime, end: datetime, branches: set[str], f: Findings) -> None:
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next((i for i, r in enumerate(rows) if r and "Occurrence Number" in r), None)
    if header_idx is None:
        f.block(f"{path.name}: no header row containing 'Occurrence Number' - is this the Origami Incidents export?")
        return
    header = [str(h) if h is not None else "" for h in rows[header_idx]]
    missing = [c for c in REQUIRED_INCIDENT_COLUMNS if c not in header]
    if missing:
        f.block(f"{path.name}: export is missing columns {missing}. Re-pull the Origami view with these columns added.")
        return
    col = {h: i for i, h in enumerate(header)}
    data = [r for r in rows[header_idx + 1:] if r and r[col["Occurrence Number"]] not in (None, "")]
    if not data:
        f.block(f"{path.name}: no incident rows.")
        return
    loss = [r[col["Loss Date"]] for r in data if isinstance(r[col["Loss Date"]], datetime)]
    outside = [d for d in loss if not (start <= d <= end)]
    f.note(f"{path.name}: {len(data)} incidents, loss dates {min(loss).date()} to {max(loss).date()}")
    if outside:
        f.block(f"{path.name}: {len(outside)} of {len(loss)} loss dates are outside {start.date()}-{end.date()}. "
                f"The export was not date-filtered to the quarter; re-pull it (or pass --start/--end to the check script knowingly).")
    locations = Counter(str(r[col["Location"]] or "").split(" ")[0] for r in data)
    for b in sorted(branches):
        code = BRANCH_LOCATION_CODE.get(b)
        if code and code not in locations:
            f.block(f"{path.name}: no incidents for location {code} ({b}). The Origami filter is probably missing that branch - "
                    f"re-pull with {code} included. (Locations present: {dict(locations)})")
    unmapped = [loc for loc in locations if loc not in BRANCH_LOCATION_CODE.values()]
    if unmapped:
        f.note(f"{path.name}: locations with no bonus workbook (ignored): {unmapped}")
    blank_ids = sum(1 for r in data if r[col["Employee Number"]] in (None, ""))
    if blank_ids / len(data) > BLANK_EMPLOYEE_NUMBER_WARN_SHARE:
        f.warn(f"{path.name}: {blank_ids} of {len(data)} incidents have no Employee Number - those are matched by name, which is less reliable. Ask Origami users to fill it in.")
    prev = Counter(str(r[col["Preventable?"]] or "blank") for r in data)
    disc = Counter(str(r[col["Disciplinary action?"]] or "blank") for r in data)
    f.note(f"{path.name}: Preventable? {dict(prev)}; Disciplinary action? {dict(disc)}")


def check_active_roster(path: Path, f: "Findings") -> None:
    """The roster must identify employees by ID and should carry a job title so Yard
    Jockeys (never eligible) can be excluded from the sheets."""
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    header_row, id_col, name_col, title_col = find_roster_columns(ws)
    if id_col is None:
        f.block(f"{path.name}: no Employee ID column found - cannot match the roster to the sheets")
        return
    if title_col is None:
        f.warn(f"{path.name}: no Job Title / Position column - Yard Jockeys cannot be excluded by title; "
               "ask Jeremy for a roster export that includes the title")
        return
    excluded = sum(1 for row in ws.iter_rows(min_row=header_row + 1, values_only=True)
                   if title_col < len(row) and is_excluded_title(row[title_col]))
    f.note(f"{path.name}: roster OK (ID + title columns found); {excluded} Yard Jockey / excluded-title rows will be kept off the sheets")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--term-report", required=True, type=Path)
    parser.add_argument("--incidents", type=Path)
    parser.add_argument("--active-drivers", type=Path, help="Latest HR active-driver list (xlsx). Ask Jeremy for it before running.")
    parser.add_argument("--start", required=True, type=lambda v: datetime.strptime(v, "%Y-%m-%d"))
    parser.add_argument("--end", required=True, type=lambda v: datetime.strptime(v, "%Y-%m-%d"))
    args = parser.parse_args()

    f = Findings()
    branches = check_branch_workbooks(args.input_dir, f)
    check_term_report(args.term_report, args.start, args.end, f)
    if args.incidents:
        check_incidents(args.incidents, args.start, args.end, branches, f)
    else:
        f.warn("No Origami incident export given - step 2 (incident cross-check) will be skipped.")
    if args.active_drivers:
        if not args.active_drivers.exists():
            f.block(f"Active-driver list not found: {args.active_drivers}")
        else:
            check_active_roster(args.active_drivers, f)
    else:
        f.warn("No ACTIVE-DRIVER list given. Ask Jeremy for the latest HR active-driver export before running - "
               "without it, drivers active in HR but missing from the sheets (and sheet rows for people no longer active) cannot be checked.")

    print(f"PREFLIGHT for {args.start.date()} to {args.end.date()}")
    for label, items in (("BLOCKERS", f.blockers), ("WARNINGS", f.warnings), ("INFO", f.info)):
        print(f"\n{label}: {len(items)}")
        for msg in items:
            print(f"  - {msg}")
    if f.blockers:
        print("\nSTOP: fix the blockers (or get Jeremy's explicit go-ahead) before running the pipeline.")
        sys.exit(1)
    print("\nOK to proceed. Tell Jeremy about the warnings before running.")


if __name__ == "__main__":
    main()
