#!/usr/bin/env python3
"""Extract .docx / .xlsx / .pptx into review-friendly text.

Grok (and any reviewer) critiques structure, claims, formulas, and logic far
better from text than from a binary file. This emits markdown-ish text:
  - docx: headings/paragraphs in order, tables as pipe rows
  - xlsx: each sheet as a bounded cell grid (values; formulas if present)
  - pptx: each slide's text and notes

Usage: python3 extract_office.py <file> [> packet_body.txt]
"""
import sys
from pathlib import Path

MAX_ROWS = 200   # bound spreadsheet output so the packet stays lean
MAX_COLS = 40


def extract_docx(path: Path) -> str:
    from docx import Document  # python-docx
    doc = Document(str(path))
    out = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower()
        if "heading 1" in style:
            out.append(f"# {text}")
        elif "heading 2" in style:
            out.append(f"## {text}")
        elif "heading 3" in style:
            out.append(f"### {text}")
        else:
            out.append(text)
    for ti, table in enumerate(doc.tables):
        out.append(f"\n[table {ti + 1}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out)


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    # data_only=False keeps formulas visible so logic can be reviewed.
    wb = load_workbook(str(path), data_only=False)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n## sheet: {ws.title}  ({ws.max_row}x{ws.max_column})")
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if r > MAX_ROWS:
                out.append(f"... ({ws.max_row - MAX_ROWS} more rows truncated)")
                break
            vals = ["" if v is None else str(v) for v in row[:MAX_COLS]]
            if any(vals):
                out.append("| " + " | ".join(vals) + " |")
    return "\n".join(out)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        out.append(f"\n## slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    out.append(t)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"[notes] {notes}")
    return "\n".join(out)


EXTRACTORS = {".docx": extract_docx, ".xlsx": extract_xlsx, ".pptx": extract_pptx}
DEPS = {".docx": "python-docx", ".xlsx": "openpyxl", ".pptx": "python-pptx"}


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: python3 extract_office.py <file>", file=sys.stderr)
        return 1
    path = Path(sys.argv[1])
    ext = path.suffix.lower()
    fn = EXTRACTORS.get(ext)
    if not fn:
        print(f"unsupported file type: {ext}", file=sys.stderr)
        return 1
    try:
        print(fn(path))
    except ImportError:
        print(f"missing dependency for {ext}: pip install {DEPS[ext]}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
